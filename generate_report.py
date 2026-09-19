import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_document(output_docx_path):
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("EtherX Innovations | Technical Deliverables & Architecture Report")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential — Lead Engineering Review — EtherXMeet & Pragna")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(120, 120, 120)

    # Styles setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("Engineering Deliverables & Architecture Report")
    t_run.font.size = Pt(22)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(15, 23, 42) # Slate-900

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    s_run = sub_p.add_run("Comprehensive Technical Review: EtherXMeet (NxtMeet) & Pragna AI Platform")
    s_run.font.size = Pt(13)
    s_run.font.color.rgb = RGBColor(71, 85, 105) # Slate-600

    # Meta Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_data = [
        [("Lead Author / Contributor:", "Vinay G K (Collaborator: Claude)"), ("Projects Covered:", "1. EtherXMeet (NxtMeet)\n2. Pragna Multilingual AI Assistant")],
        [("Review Focus:", "Smart Contracts, WebRTC, AI Agents, Document Gen, DB & Security"), ("Document Status:", "Complete & Production-Verified")]
    ]
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(3.4)
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            lbl, val = meta_data[row_idx][col_idx]
            r_lbl = p.add_run(lbl + " ")
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.color.rgb = RGBColor(15, 23, 42)
            r_val = p.add_run(val)
            r_val.font.size = Pt(9.5)
            r_val.font.color.rgb = RGBColor(51, 65, 85)
    set_table_borders(meta_table, color="E2E8F0", sz="4")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper for adding sections
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate-900
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate-800
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(bold_prefix)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)
        r2 = p.add_run(" " + text)
        r2.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # =========================================================================
    # PROJECT 1: ETHERXMEET (NXTMEET)
    # =========================================================================
    add_h1("Project 1: EtherXMeet (NxtMeet) — Web3 Video Conferencing")

    add_h2("1. Smart Contracts & Web3 Infrastructure (Solidity & Hardhat)")
    add_bullet("MeetingRegistry.sol:", "Implemented complete decentralized meeting lifecycle on Polygon. Handles createMeeting(meetingId, metadataCID) with IPFS metadata, joinMeeting(meetingId) for on-chain attendance logging, sendMessage(meetingId, contentCID) anchoring chat messages via zero-storage events, and endMeeting(meetingId, notesHash) sealing keccak256 AI meeting summaries.")
    add_bullet("MeetingNFT.sol (ERC721URIStorage):", "Implemented an NFT meeting receipt / POAP minter linked to MeetingRegistry, allowing hosts to mint verifiable ERC-721 tokens pointing to ipfs://<cid> upon session closure.")
    add_bullet("Hardhat Tooling & Automated Deployment:", "Built multi-network deployment scripts (deploy.js) targeting Localhost, Polygon Amoy Testnet, and Polygon Mainnet, backed by a comprehensive unit/integration test suite (MeetingRegistry.test.js).")

    add_h2("2. Real-Time Video, WebRTC & In-Room Feature Suite")
    add_bullet("WebRTC Mesh Engine (useWebRTC.js):", "Custom peer-to-peer connection management covering SDP offer/answer negotiations, trickle ICE candidate buffering, dynamic track renegotiation, and real-time device toggles.")
    add_bullet("Command Center (VideoRoom.jsx):", "Grid and filmstrip layouts, active speaker spotlighting, pin/fullscreen modes, client-side screen recording (MediaRecorder API), bookmarks, and unified side-panel docking.")
    add_bullet("Live Speech Transcription (LiveTranscript.jsx):", "Continuous client-side STT via Web Speech API with rolling speaker transcripts, timestamp tracking, auto-scroll with lock, and .txt export.")
    add_bullet("Procedural Ambient Sound Mixer (AmbientSoundMixer.jsx):", "8-soundscape audio generator (Rain, Coffee Shop, Forest, Ocean, Fireplace, Lo-fi Beats, Office Hum, Thunderstorm) synthesized entirely in-browser via Web Audio API (pink/brown noise, biquad filters, LFO modulators; zero external assets).")
    add_bullet("Smart Agenda Timer (AgendaTimer.jsx):", "Multi-item structured meeting agenda with countdown timers (2m-30m presets), overtime alerts, and visual color-coded progress bars.")
    add_bullet("Live Analytics Dashboard (MeetingAnalytics.jsx):", "Real-time meeting health metrics, participant talk-time breakdown, dynamic engagement scoring, and word frequency clouds.")
    add_bullet("Collaborative Whiteboard (Whiteboard.jsx):", "Interactive canvas supporting vector shapes, freehand drawing, color palettes, and sub-millisecond Socket.IO multi-user sync.")

    add_h2("3. Embedded Wallet & Gasless Onboarding")
    add_bullet("Embedded Ethers Wallet (WalletContext.jsx):", "Auto-generated client-side ethers.js wallet initialized upon user login, storing encrypted keys in localStorage to eliminate mandatory MetaMask friction for mainstream users.")
    add_bullet("Operator Drip Faucet (POST /api/wallet/fund):", "Automated backend operator wallet providing gas subsidies (MATIC/POL) to newly created user addresses on Polygon Amoy.")
    add_bullet("Frictionless Hybrid Access:", "Enabled direct room entry for guest participants while preserving optional on-chain receipts, token gating (TokenGateModal.jsx), and verified chat for Web3 users.")

    add_h2("4. Backend Services, Signaling & APIs")
    add_bullet("WebRTC Signaling Gateway (signaling.js):", "Socket.IO gateway managing room lifecycle, peer routing, ICE candidate exchange, chat broadcasting, whiteboard sync, and attendance records.")
    add_bullet("Auth & Recording Pipelines:", "JWT session authentication, Google OAuth with Passport, multipart .webm upload handler with blob streaming playback, and token-balance gating validation.")
    add_bullet("Deployment & LAN Bindings:", "Configured multi-origin CORS and LAN IP bindings for mobile browser testing, alongside Vercel serverless functions.")

    add_h2("5. UI/UX, Design System & Mobile Optimization")
    add_bullet("Custom Shader Visuals:", "Canvas background shaders (ShaderBackground, AuroraBackground, ParticleNetwork, GoldGlitter) creating a distinct modern aesthetic.")
    add_bullet("Mobile Stabilization:", "Resolved mobile white-screen issues with solid #0A0A0F fallbacks, single-column responsive viewports, and custom UI components.")

    # Commit Table for EtherXMeet
    add_h2("EtherXMeet Chronological Commit Milestones")
    
    commit_table = doc.add_table(rows=1, cols=4)
    commit_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    commit_table.autofit = False
    
    headers = ["Commit Hash", "Date", "Area", "Key Deliverables"]
    widths = [Inches(1.1), Inches(1.0), Inches(1.1), Inches(3.6)]
    
    hdr_cells = commit_table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "0F172A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    commits_data = [
        ("f4694f7", "May 08, 2026", "Core Architecture", "Initial commit: EtherXMeet Solidity contracts, frontend, and backend."),
        ("0e8ed09–fbeff51", "May 08, 2026", "Config & Network", "Environment variables, Polygon network configs, and script scaffolding."),
        ("68b19df", "May 08, 2026", "Bugfix & Media", "Blob-based video playback, recordings sync, and auth background transparency."),
        ("bb38122–6e4a774", "May 14, 2026", "Web3 / Wallet", "Embedded wallet design spec, User model update (walletFunded), Amoy RPC."),
        ("d37c79e–4e482f5", "May 14, 2026", "Backend Faucet", "Backend drip wallet endpoint (POST /api/wallet/fund) with ethers.js."),
        ("429d4b1–6ad0348", "May 14, 2026", "Frontend Wallet", "Embedded WalletContext replacing MetaMask-only flows; updated contract hooks."),
        ("b140296–60573c8", "May 14, 2026", "UX / Onboarding", "Wallet banner states, transaction failure skips, and direct meeting join."),
        ("a9053ec", "Jun 29, 2026", "Major Feature", "Full UI overhaul, useWebRTC hook, signaling.js, VideoRoom, Whiteboard."),
        ("e0d2f69–f84b276", "Jun 30, 2026", "In-Room Audio/STT", "Live speech transcription (Web Speech API) & synthesized ambient sound mixer."),
        ("77b554a–5e780a8", "Jun 30, 2026", "Analytics & Timer", "Live meeting analytics dashboard and Smart Agenda countdown timer."),
        ("6fe3164–dd0503f", "Jun 30, 2026", "Integration & Fix", "Integrated side-panels in VideoRoom, mobile white screen fix, LAN CORS."),
        ("7d9946e–c61c2c1", "Aug 30, 2026", "Branding & Deploy", "Brand transition to NxtMeet, QR modal, Vercel serverless integration.")
    ]
    
    for row_idx, data in enumerate(commits_data):
        row = commit_table.add_row()
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for col_idx, cell in enumerate(row.cells):
            cell.width = widths[col_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(data[col_idx])
            run.font.size = Pt(8.5)
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(15, 23, 42)
            else:
                run.font.color.rgb = RGBColor(51, 65, 85)
                
    set_table_borders(commit_table, color="E2E8F0", sz="4")

    # =========================================================================
    # PROJECT 2: PRAGNA AI CHATBOT & AGENTIC PLATFORM
    # =========================================================================
    doc.add_page_break()
    add_h1("Project 2: Pragna — Multilingual AI Chatbot & Agentic Platform")

    add_h2("1. Autonomous Coding Agent (Web & CLI Engine)")
    add_bullet("Dual-Surface Execution Loop:", "Built a think-tool-observe autonomous execution loop available both as an in-browser streaming panel (via SSE) and a standalone terminal CLI (pragna_code.py).")
    add_bullet("Filesystem Sandboxing:", "Enforced strict directory boundaries with _resolve_in_root() across all file tools (read_file, write_file, create_file, append_file, list_dir, search_code), blocking path traversal attacks.")
    add_bullet("Human-in-the-Loop Approval Gate:", "Separated safe read-only tools from mutating tools (write_file, run_command). Mutating actions pause execution, generating visual diffs and command previews requiring explicit user approval.")
    add_bullet("Resilient Agent Sessions:", "Snapshotting in-memory AGENT_SESSIONS before eviction cycles, resolving concurrency race conditions.")

    add_h2("2. Multi-Format AI Document Generation Engine")
    add_bullet("Native Document Generation:", "Built automated generation pipelines producing styled Microsoft Word (.docx), PDF (.pdf), PowerPoint (.pptx), and Excel (.xlsx) files.")
    add_bullet("Markdown Outline Parsing:", "Engineered an outline parser converting structured LLM outputs into branded documents featuring cover pages, headers, data tables, and slide decks.")
    add_bullet("Conversational Integration:", "Automated intent classification detecting document generation prompts across input bars and quick prompts, delivering downloadable attachment cards directly inside chat.")

    add_h2("3. Enterprise Chat UX & Knowledge Management")
    add_bullet("Real-Time Token Streaming:", "Configured Server-Sent Events (SSE) across /api/chat_stream, unifying all chat send triggers (InputBar, suggestions, retries) into instant streaming responses.")
    add_bullet("Conversation Management Suite:", "Implemented chat branching/duplication, collapsible folder hierarchies, message bookmarking, in-line editing/regeneration, and multi-field transcript search.")
    add_bullet("Summarization & Templates:", "Engineered /api/summarize_chat for instant conversation digests and built a customizable Prompt Template Library with full CRUD.")
    add_bullet("Custom Personas & Slash Commands:", "Built an interactive slash-command menu (/image, /doc, /persona) with autocomplete, paired with dynamic system prompt injection and database-backed persona CRUD.")

    add_h2("4. Backend Infrastructure & Database Security")
    add_bullet("Database Modernization:", "Migrated from local SQLite to persistent Supabase PostgreSQL with automated schema initializers.")
    add_bullet("Connection Pooling (psycopg_pool):", "Built per-process connection pool allocation preventing socket collisions in forked Gunicorn workers. Added pre-checkout liveness tests (SELECT 1) and health metrics on /api/health.")
    add_bullet("Multi-Tier LLM Gateway:", "Configured resilient provider fallback tiers (Ollama local/cloud -> Groq -> OpenAI) with bounded LRU caches to prevent OOM termination.")
    add_bullet("Authentication & Transactional Email:", "Implemented JWT security, mandatory email OTP signup verification, background password resets (EmailJS/SMTP/Resend), and Google/GitHub OAuth 2.0 flows.")

    add_h2("5. Frontend & Visual Polish")
    add_bullet("Dark-Gold Noir Theme:", "Rebuilt styling into a cohesive pure black (#000000) and dark-gold design with responsive mobile, tablet, and desktop views.")
    add_bullet("Live HTML Artifacts:", "Built a sandboxed iframe side panel rendering interactive web applications, code viewers, and download tools in real-time.")
    add_bullet("Visuals & Splash:", "Created an interactive WebGL neural vortex shader canvas and regional language splash sequences.")

    # Pragna Milestones Table
    add_h2("Pragna Chronological Commit Milestones")
    
    pragna_table = doc.add_table(rows=1, cols=4)
    pragna_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pragna_table.autofit = False
    
    hdr_cells2 = pragna_table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells2[i].width = widths[i]
        set_cell_background(hdr_cells2[i], "0F172A")
        set_cell_margins(hdr_cells2[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells2[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    pragna_data = [
        ("8b2df19–87784d3", "Jul 02, 2026", "LLM Gateway", "Configured Ollama primary models with cloud fallbacks; fixed prompt handling."),
        ("bca881b–6f98085", "Jul 04, 2026", "Agent Sandbox", "Built code agent loop, directory sandboxing, confirm-before-act gates & diff previews."),
        ("08e9b88–24bb43b", "Jul 05, 2026", "Chat Folders/Edit", "Added chat folders, in-line message editing, markdown export, and sidebar search."),
        ("0a257b9–6488426", "Jul 06, 2026", "Shortcuts/Branches", "Implemented message bookmarks, chat duplicate/branching, and shortcuts overlay."),
        ("f184cfd–24385a7", "Jul 08, 2026", "RAG & Citations", "Added RAG scheduler view, model profile picker, PDF export, and source citations."),
        ("aea2fb6–5e27ab8", "Jul 10, 2026", "SSE Streaming", "Unified SSE chunked streaming across input paths; added /api/summarize_chat."),
        ("6c39ed1–213d93a", "Jul 11, 2026", "Document Engine", "Built docx/pdf/pptx/xlsx builders, outline parser, and chat attachment rendering."),
        ("a2a5f35–9797c4d", "Jul 13, 2026", "Slash Commands", "Added custom Personas CRUD and /image, /doc, /persona slash command autocomplete."),
        ("77d90aa–65e4ce0", "Jul 14–16, 2026", "Docker & Deploy", "Docker containerization, Render deploy config, and memory cache bounding."),
        ("da4ced1–76fd7af", "Jul 17–Aug 01", "Supabase & Auth", "Migrated to Supabase Postgres, psycopg_pool, email OTP verification, and SMTP."),
        ("acc0963–df66e91", "Aug 03–08, 2026", "OAuth & Noir UI", "Added Google/GitHub OAuth 2.0, WebGL neural vortex, and pure black styling."),
        ("8e2f78d–9d708bb", "Aug 16–Sep 07", "Live Artifacts", "Added standalone HTML Live Artifacts panel, Web Speech TTS chunking, and OCR.")
    ]
    
    for row_idx, data in enumerate(pragna_data):
        row = pragna_table.add_row()
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for col_idx, cell in enumerate(row.cells):
            cell.width = widths[col_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(data[col_idx])
            run.font.size = Pt(8.5)
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(15, 23, 42)
            else:
                run.font.color.rgb = RGBColor(51, 65, 85)
                
    set_table_borders(pragna_table, color="E2E8F0", sz="4")

    # =========================================================================
    # COMPARATIVE SUMMARY TABLE
    # =========================================================================
    doc.add_page_break()
    add_h1("Cross-Project Architectural Matrix")
    
    summary_table = doc.add_table(rows=1, cols=3)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_table.autofit = False
    
    s_headers = ["Architectural Dimension", "EtherXMeet (NxtMeet)", "Pragna AI Platform"]
    s_widths = [Inches(1.8), Inches(2.5), Inches(2.5)]
    
    s_hdr_cells = summary_table.rows[0].cells
    for i, h in enumerate(s_headers):
        s_hdr_cells[i].width = s_widths[i]
        set_cell_background(s_hdr_cells[i], "0F172A")
        set_cell_margins(s_hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = s_hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    s_data = [
        ("Core Capability", "P2P WebRTC Video Conferencing & Web3 Meeting Verification", "Multilingual LLM Assistant, Autonomous Coding Agent & Document Gen"),
        ("Smart Contracts / Web3", "Solidity, Polygon Amoy/Mainnet, Hardhat, MeetingRegistry, MeetingNFT (ERC-721)", "Optional Web3 integration; focuses on LLM orchestration and data processing"),
        ("Real-Time & Audio", "Custom WebRTC mesh, Web Audio API procedural sound mixer, Web Speech STT", "SSE token streaming, Web Speech TTS chunking, WebGL shader animation"),
        ("File & Media Ops", "Multipart .webm recording uploads, blob streaming playback, canvas whiteboard", "Native .docx, .pdf, .pptx, .xlsx generation, live HTML Artifacts sandbox"),
        ("Backend & Database", "Node.js / Express, Socket.IO, MongoDB / Mongoose models, JWT & OAuth", "Python Flask, Supabase PostgreSQL, psycopg_pool, FAISS RAG, JWT & OAuth"),
        ("Security & Sandboxing", "Token-gated rooms, operator faucet rate limits, verified chat signatures", "Filesystem directory traversal sandboxing, confirm-before-act tool approval gate"),
        ("Onboarding UX", "Client-side embedded Ethers wallet (gasless) + frictionless guest join", "Email OTP signup verification, background password resets, OAuth 2.0"),
        ("Frontend & Styling", "React / Vite, custom canvas shaders (Aurora, Glitter), responsive dark mode", "React / Vite, dark-gold noir (#000000), interactive WebGL vortex, responsive UI")
    ]
    
    for row_idx, data in enumerate(s_data):
        row = summary_table.add_row()
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for col_idx, cell in enumerate(row.cells):
            cell.width = s_widths[col_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(data[col_idx])
            run.font.size = Pt(8.5)
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(15, 23, 42)
            else:
                run.font.color.rgb = RGBColor(51, 65, 85)
                
    set_table_borders(summary_table, color="E2E8F0", sz="4")

    # Conclusion paragraph
    c_p = doc.add_paragraph()
    c_p.paragraph_format.space_before = Pt(16)
    c_p.paragraph_format.space_after = Pt(4)
    c_run = c_p.add_run("Summary Assessment:")
    c_run.font.bold = True
    c_run.font.size = Pt(11)
    c_run.font.color.rgb = RGBColor(15, 23, 42)
    
    c_desc = doc.add_paragraph()
    c_desc.paragraph_format.space_after = Pt(4)
    c_desc_run = c_desc.add_run(
        "Both repositories represent end-to-end engineered, production-ready systems. EtherXMeet pioneers frictionless Web3 video collaboration with client-side synthesis and smart contract verification, while Pragna delivers an enterprise-grade AI assistant with strict sandbox security, document generation, and a resilient database architecture."
    )
    c_desc_run.font.size = Pt(10)
    c_desc_run.font.color.rgb = RGBColor(51, 65, 85)

    doc.save(output_docx_path)
    print(f"Successfully generated docx at: {output_docx_path}")

if __name__ == "__main__":
    out_path = sys.argv[1] if len(sys.argv) > 1 else "/home/vinay/EtherX_Projects_Executive_Report.docx"
    build_document(out_path)
